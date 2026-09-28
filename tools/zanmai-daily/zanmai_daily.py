#!/usr/bin/env python3
"""
毎日 8:00 に Windows のタスクスケジューラから実行する想定のスクリプト。

1. Excel（シート「API」）の D 列から、まだページに載っていない値を最大 100 個取り出す
2. サーバ上の index.html を FTP でダウンロードし、ローカルにバックアップする
3. STAGE START〜STAGE END の <ul class="video-row"> の末尾に追記する
4. 更新した index.html を FTP でサーバへアップロードする

使い方:
  py zanmai_daily.py            本番実行
  py zanmai_daily.py --dry-run  アップロードせず preview.html だけ作る
  py zanmai_daily.py --force    今日すでに実行済みでももう一度実行する
"""
import argparse
import datetime as dt
import ftplib
import html
import io
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
JST = dt.timezone(dt.timedelta(hours=9))
NOW = dt.datetime.now(JST)


def log(msg):
    line = f"[{NOW:%Y-%m-%d %H:%M:%S}] {msg}"
    print(line)
    (HERE / "logs").mkdir(exist_ok=True)
    with open(HERE / "logs" / f"{NOW:%Y-%m}.log", "a", encoding="utf-8") as f:
        f.write(line + "\n")


def load_config():
    path = HERE / "config.json"
    if not path.exists():
        sys.exit("config.json がありません。config.example.json をコピーして作ってください。")
    return json.loads(path.read_text(encoding="utf-8"))


# ---------- Excel ----------

def find_excel(path_str):
    p = Path(path_str)
    if p.is_file():
        return p
    if p.is_dir():
        files = [f for f in p.iterdir()
                 if f.suffix.lower() in (".xlsx", ".xlsm") and not f.name.startswith("~$")]
        if len(files) == 1:
            return files[0]
        names = ", ".join(f.name for f in files) or "（なし）"
        sys.exit(f"{p} の Excel ファイルを 1 つに特定できません: {names}\n"
                 "config.json の excel_path にファイル名まで書いてください。")
    sys.exit(f"excel_path が見つかりません: {p}")


def read_column(cfg):
    from openpyxl import load_workbook  # pip install openpyxl

    src = find_excel(cfg["excel_path"])
    # Excel で開いたままでも読めるよう、一時フォルダにコピーしてから読む
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / src.name
        shutil.copy2(src, copy)
        wb = load_workbook(copy, read_only=True, data_only=True)
        if cfg["sheet"] not in wb.sheetnames:
            sys.exit(f"シート「{cfg['sheet']}」がありません。存在するシート: {wb.sheetnames}")
        ws = wb[cfg["sheet"]]
        col = cfg["column"]
        values = []
        for (cell,) in ws.iter_rows(min_row=cfg.get("start_row", 2),
                                    min_col=_col_index(col), max_col=_col_index(col)):
            v = cell.value
            if v is None or str(v).strip() == "":
                continue
            values.append(str(v).strip())
        wb.close()
    log(f"Excel 読み込み: {src.name} / {cfg['sheet']}!{col} → {len(values)} 件")
    return values


def _col_index(letter):
    n = 0
    for ch in letter.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


# ---------- HTML ----------

CID_RE = re.compile(r"cid=([A-Za-z0-9_]+)")


def dedupe_key(value):
    """同じ作品を二重に載せないための判定キー。cid があれば cid、なければ値そのもの。"""
    m = CID_RE.search(value)
    if m:
        return "cid=" + m.group(1)
    if re.fullmatch(r"[A-Za-z0-9_]+", value):  # D 列が cid だけの場合
        return "cid=" + value
    return value


def render(value, cfg):
    if value.lstrip().startswith("<"):
        return value  # D 列が HTML ならそのまま入れる
    return cfg["item_template"].replace("{value}", value).replace("{value_escaped}", html.escape(value))


def insert_items(page, items_html, cfg):
    start = page.find(cfg["start_marker"])
    end = page.find(cfg["end_marker"])
    if start == -1 or end == -1 or end < start:
        sys.exit(f"HTML に「{cfg['start_marker']}」「{cfg['end_marker']}」が見つかりません。中断します。")
    close_ul = page.rfind("</ul>", start, end)
    if close_ul == -1:
        sys.exit("STAGE 区間の中に </ul> が見つかりません。中断します。")
    block = f"\n      <!-- ---------- 追加 {NOW:%Y-%m-%d} ---------- -->\n" + "\n".join(items_html) + "\n\n    "
    # </ul> 直前の空白を整えてから差し込む
    head = page[:close_ul].rstrip() + "\n"
    return head + block + page[close_ul:]


# ---------- FTP ----------

def ftp_connect(cfg):
    f = cfg["ftp"]
    cls = ftplib.FTP_TLS if f.get("protocol", "ftps") == "ftps" else ftplib.FTP
    ftp = cls()
    ftp.encoding = "utf-8"
    ftp.connect(f["host"], f.get("port", 21), timeout=60)
    ftp.login(f["user"], f["password"])
    if isinstance(ftp, ftplib.FTP_TLS):
        ftp.prot_p()
    return ftp


def ftp_download(ftp, remote):
    buf = io.BytesIO()
    ftp.retrbinary(f"RETR {remote}", buf.write)
    return buf.getvalue().decode("utf-8")


def ftp_upload(ftp, remote, text):
    ftp.storbinary(f"STOR {remote}", io.BytesIO(text.encode("utf-8")))


# ---------- main ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    cfg = load_config()
    state_path = HERE / "state.json"
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    today = f"{NOW:%Y-%m-%d}"
    if state.get("last_run") == today and not (args.force or args.dry_run):
        log(f"本日（{today}）は実行済みのためスキップ（やり直すときは --force）")
        return

    values = read_column(cfg)

    ftp = ftp_connect(cfg)
    try:
        remote = cfg["ftp"]["remote_html"]
        page = ftp_download(ftp, remote)
        backups = HERE / "backups"
        backups.mkdir(exist_ok=True)
        (backups / f"index_{NOW:%Y%m%d_%H%M%S}.html").write_text(page, encoding="utf-8")
        for old in sorted(backups.glob("index_*.html"))[:-cfg.get("keep_backups", 30)]:
            old.unlink()

        stage = page[page.find(cfg["start_marker"]):page.find(cfg["end_marker"])]
        seen = set()
        picked = []
        for v in values:
            key = dedupe_key(v)
            if key in seen or key in stage:
                continue
            seen.add(key)
            picked.append(v)
            if len(picked) >= cfg.get("count", 100):
                break

        if not picked:
            log("追加できる新しい値がありません（すべて掲載済み）。終了します。")
            return
        if len(picked) < cfg.get("count", 100):
            log(f"注意: 新しい値が {len(picked)} 件しかありませんでした")

        new_page = insert_items(page, [render(v, cfg) for v in picked], cfg)

        if args.dry_run:
            (HERE / "preview.html").write_text(new_page, encoding="utf-8")
            log(f"[dry-run] {len(picked)} 件を追記した preview.html を作成（アップロードはしていません）")
            return

        ftp_upload(ftp, remote, new_page)
        if cfg.get("local_copy"):
            Path(cfg["local_copy"]).write_text(new_page, encoding="utf-8")
    finally:
        try:
            ftp.quit()
        except Exception:
            ftp.close()

    state.update(last_run=today, last_added=len(picked))
    state_path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    log(f"完了: {len(picked)} 件を追記して {remote} にアップロードしました")


if __name__ == "__main__":
    try:
        main()
    except SystemExit as e:
        if e.code not in (None, 0):
            log(f"エラー: {e.code}")
        raise
    except Exception as e:
        log(f"エラー: {type(e).__name__}: {e}")
        raise
